#!/usr/bin/env python3
"""Audit and transform the frozen China hourly carbon-intensity source.

This is a zero-search source gate.  It verifies the Figshare v3 metadata and
workbook, profiles every scenario year, converts S1-2025 hourly values to the
paper's 48 half-hour slots without interpolation, and pre-registers Shanghai
representative days from the carbon curve alone.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import platform
import sys
from calendar import isleap
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[2]
RAW = REPO / "data/Carbon/中国情景/raw_20260717"
DEFAULT_OUTPUT = REPO / "baselines/e4_e5/china_tvci_source_gate_20260717_v2"
S1 = RAW / "CEF_data_Scenario_S1.xlsx"
FIGSHARE_METADATA = RAW / "figshare_article_28953545.json"
SOURCES = {
    "CEF_data_Scenario_S1.xlsx": "8e23df4a39707077a2c772e9ab0a275b388fe4ded04d0db4fe0937616499486d",
    "figshare_article_28953545.json": "c4729e9cd5d2de21652b12185729c560d04d454cb6d43ac7888ffde181c936d5",
    "figshare_dataset_annotation.pdf": "75d3d0a265cc3e1df044f00c8a10f4328f85c9de776582df4f22b2a75916691d",
    "MEE_2023_power_CO2_factors.pdf": "d43b60a3eaad9f3fe59ca37c204d3d0e1b6ea52742a092f0d856971a3f018e5a",
    "Shanghai_2025_07_industrial_tou_tariff_10kV.pdf": "e2293dcc5fb15da3eaad204c929447a1a6dfa757248610b38ef1dec674157bf9",
    "NDRC_land_transport_GHG_guideline.pdf": "84aadb948141b5229f9339254dd305bbf54b435e140827244d9e3c69b289bb31",
    "MEE_2023_power_carbon_footprint_factors.pdf": "68afcb5a56d63311642fa4974004a13e87c693d3de16d934f718977af080ad09",
}
EXPECTED_MD5 = "3cdf56af1e9c90167fc4c85ccc5afd5b"
EXPECTED_YEARS = tuple(str(year) for year in range(2025, 2061, 5))
EXPECTED_COLUMNS = (
    "Time",
    "Mainland China",
    "Beijing",
    "Tianjin",
    "Hebei",
    "Shanxi",
    "Inner Mongolia",
    "Liaoning",
    "Jilin",
    "Heilongjiang",
    "Shanghai",
    "Jiangsu",
    "Zhejiang",
    "Anhui",
    "Fujian",
    "Jiangxi",
    "Shandong",
    "Henan",
    "Hubei",
    "Hunan",
    "Guangdong",
    "Guangxi",
    "Hainan",
    "Chongqing",
    "Sichuan",
    "Guizhou",
    "Yunnan",
    "Tibet",
    "Shaanxi",
    "Gansu",
    "Qinghai",
    "Ningxia",
    "Xinjiang",
)
MEE_2023 = {
    "Mainland China": 0.5306,
    "Shanghai": 0.5737,
    "Jiangsu": 0.5827,
}


class SourceGateError(RuntimeError):
    """Raised when a frozen source violates the pre-registered contract."""


def sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def md5_path(path: Path) -> str:
    digest = hashlib.md5(usedforsecurity=False)
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        handle.write(content)
        handle.flush()
        os.fsync(handle.fileno())


def write_json(path: Path, payload: Any) -> None:
    write_text(
        path,
        json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)
        + "\n",
    )


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise SourceGateError(f"refuse to write empty CSV: {path.name}")
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    write_text(path, buffer.getvalue())


def source_checks() -> tuple[dict[str, dict[str, Any]], list[str]]:
    records: dict[str, dict[str, Any]] = {}
    failures: list[str] = []
    for name, expected in SOURCES.items():
        path = RAW / name
        if not path.is_file():
            failures.append(f"missing frozen source: {name}")
            continue
        actual = sha256_path(path)
        records[name] = {
            "path": str(path.relative_to(REPO)),
            "bytes": path.stat().st_size,
            "sha256": actual,
            "expected_sha256": expected,
            "hash_match": actual == expected,
        }
        if actual != expected:
            failures.append(f"source hash differs: {name}")
    if S1.is_file() and md5_path(S1) != EXPECTED_MD5:
        failures.append("Figshare S1 MD5 differs")
    return records, failures


def verify_figshare_metadata() -> tuple[dict[str, Any], list[str]]:
    failures: list[str] = []
    payload = json.loads(FIGSHARE_METADATA.read_text(encoding="utf-8"))
    if payload.get("id") != 28953545 or payload.get("version") != 3:
        failures.append("Figshare article identity/version differs")
    if payload.get("doi") != "10.6084/m9.figshare.28953545.v3":
        failures.append("Figshare DOI differs")
    license_payload = payload.get("license", {})
    if license_payload.get("name") != "CC BY 4.0":
        failures.append("Figshare license differs")
    files = {int(row["id"]): row for row in payload.get("files", [])}
    s1 = files.get(63393795, {})
    if s1.get("name") != "CEF data for Scenario S1.xlsx":
        failures.append("Figshare S1 filename differs")
    if s1.get("computed_md5") != EXPECTED_MD5 or int(s1.get("size", -1)) != S1.stat().st_size:
        failures.append("Figshare S1 metadata size/MD5 differs")
    return {
        "article_id": payload.get("id"),
        "version": payload.get("version"),
        "doi": payload.get("doi"),
        "title": payload.get("title"),
        "license": license_payload,
        "published_date": payload.get("published_date"),
        "modified_date": payload.get("modified_date"),
        "s1_file": s1,
    }, failures


def expected_hours(year: int) -> int:
    return 8784 if isleap(year) else 8760


def audit_workbook() -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[tuple[Any, ...]],
    list[str],
    list[str],
]:
    workbook = load_workbook(S1, read_only=True, data_only=True)
    failures: list[str] = []
    warnings: list[str] = []
    if tuple(workbook.sheetnames) != EXPECTED_YEARS:
        failures.append(f"sheet sequence differs: {workbook.sheetnames}")
    year_rows: list[dict[str, Any]] = []
    summary_rows: list[dict[str, Any]] = []
    values_2025: list[tuple[Any, ...]] = []
    for sheet_name in workbook.sheetnames:
        year = int(sheet_name)
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = tuple(next(rows))
        column_match = header == EXPECTED_COLUMNS
        if not column_match:
            failures.append(f"{year}: column sequence differs")
        expected = expected_hours(year)
        count = 0
        invalid_time = 0
        null_count = 0
        non_numeric_count = 0
        negative_count = 0
        nonfinite_count = 0
        above_two_count = 0
        sums = [0.0] * (len(EXPECTED_COLUMNS) - 1)
        valid_counts = [0] * (len(EXPECTED_COLUMNS) - 1)
        missing_counts = [0] * (len(EXPECTED_COLUMNS) - 1)
        minima = [math.inf] * (len(EXPECTED_COLUMNS) - 1)
        maxima = [-math.inf] * (len(EXPECTED_COLUMNS) - 1)
        for count, row in enumerate(rows, start=1):
            if len(row) != len(EXPECTED_COLUMNS):
                failures.append(f"{year}: row {count} width differs")
                continue
            if row[0] != count:
                invalid_time += 1
            if year == 2025:
                values_2025.append(tuple(row))
            for index, value in enumerate(row[1:]):
                if value is None:
                    null_count += 1
                    missing_counts[index] += 1
                    continue
                if not isinstance(value, (int, float)):
                    non_numeric_count += 1
                    continue
                number = float(value)
                if not math.isfinite(number):
                    nonfinite_count += 1
                    continue
                if number < 0:
                    negative_count += 1
                if number > 2:
                    above_two_count += 1
                sums[index] += number
                valid_counts[index] += 1
                minima[index] = min(minima[index], number)
                maxima[index] = max(maxima[index], number)
        if count != expected:
            failures.append(f"{year}: expected {expected} hourly rows, got {count}")
        for label, total, valid_count, missing_count, minimum, maximum in zip(
            EXPECTED_COLUMNS[1:],
            sums,
            valid_counts,
            missing_counts,
            minima,
            maxima,
            strict=True,
        ):
            mean = total / valid_count if valid_count else math.nan
            summary_rows.append(
                {
                    "year": year,
                    "region": label,
                    "hour_count": count,
                    "valid_value_count": valid_count,
                    "missing_value_count": missing_count,
                    "mean_tCO2_per_MWh": f"{mean:.8f}" if math.isfinite(mean) else "",
                    "min_tCO2_per_MWh": f"{minimum:.8f}" if math.isfinite(minimum) else "",
                    "max_tCO2_per_MWh": f"{maximum:.8f}" if math.isfinite(maximum) else "",
                    "MEE_2023_kgCO2_per_kWh": (
                        f"{MEE_2023[label]:.4f}" if label in MEE_2023 else ""
                    ),
                    "relative_difference_to_MEE_2023_pct": (
                        f"{100.0 * (mean / MEE_2023[label] - 1.0):.6f}"
                        if label in MEE_2023 and year == 2025
                        else ""
                    ),
                }
            )
        row_failures = (
            invalid_time
            + null_count
            + non_numeric_count
            + negative_count
            + nonfinite_count
            + above_two_count
        )
        if row_failures:
            message = f"{year}: invalid cell/time count {row_failures}"
            if year == 2025:
                failures.append(message)
            else:
                warnings.append(message + "; this unused projection year is not authorized")
        year_rows.append(
            {
                "year": year,
                "expected_hours": expected,
                "observed_hours": count,
                "column_count": len(header),
                "column_sequence_match": int(column_match),
                "invalid_time_count": invalid_time,
                "null_count": null_count,
                "non_numeric_count": non_numeric_count,
                "negative_count": negative_count,
                "nonfinite_count": nonfinite_count,
                "above_2_tCO2_per_MWh_count": above_two_count,
                "search_evaluations": 0,
                "source_gate_pass": int(count == expected and column_match and row_failures == 0),
            }
        )
    return year_rows, summary_rows, values_2025, failures, warnings


def half_hour_rows(values: list[tuple[Any, ...]]) -> tuple[list[dict[str, Any]], list[str]]:
    failures: list[str] = []
    output: list[dict[str, Any]] = []
    start = date(2025, 1, 1)
    for hour_index, row in enumerate(values):
        day_index, hour = divmod(hour_index, 24)
        current_date = start + timedelta(days=day_index)
        for half in (0, 1):
            record: dict[str, Any] = {
                "date": current_date.isoformat(),
                "day_index": day_index + 1,
                "half_hour_slot": hour * 2 + half + 1,
                "source_hour": hour_index + 1,
                "hour_of_day": hour,
                "minute": half * 30,
            }
            record.update(
                {label: f"{float(value):.8f}" for label, value in zip(EXPECTED_COLUMNS[1:], row[1:], strict=True)}
            )
            output.append(record)
    if len(output) != 365 * 48:
        failures.append(f"expected 17520 half-hour rows, got {len(output)}")
    for day_index in range(365):
        hourly = values[day_index * 24 : (day_index + 1) * 24]
        slots = output[day_index * 48 : (day_index + 1) * 48]
        for label_index, label in enumerate(EXPECTED_COLUMNS[1:], start=1):
            hourly_integral = sum(float(row[label_index]) for row in hourly)
            slot_integral = 0.5 * sum(float(row[label]) for row in slots)
            if not math.isclose(hourly_integral, slot_integral, rel_tol=0.0, abs_tol=1e-10):
                failures.append(f"day {day_index + 1} {label}: 48-slot integral differs")
    return output, failures


def select_shanghai_days(values: list[tuple[Any, ...]]) -> list[dict[str, Any]]:
    shanghai_index = EXPECTED_COLUMNS.index("Shanghai")
    start = date(2025, 1, 1)
    daily: list[dict[str, Any]] = []
    for day_index in range(365):
        day_values = [float(row[shanghai_index]) for row in values[day_index * 24 : (day_index + 1) * 24]]
        daily.append(
            {
                "date": start + timedelta(days=day_index),
                "day_index": day_index + 1,
                "mean": sum(day_values) / 24,
                "range": max(day_values) - min(day_values),
                "minimum": min(day_values),
                "maximum": max(day_values),
            }
        )
    selected: list[dict[str, Any]] = []
    quarter_bounds = ((1, 90), (91, 181), (182, 273), (274, 365))
    for quarter, (lo, hi) in enumerate(quarter_bounds, start=1):
        rows = daily[lo - 1 : hi]
        target = median(row["mean"] for row in rows)
        choices = (
            ("median_daily_mean", min(rows, key=lambda row: (abs(row["mean"] - target), row["date"]))),
            ("maximum_intraday_range", max(rows, key=lambda row: (row["range"], -row["day_index"]))),
            ("minimum_intraday_range", min(rows, key=lambda row: (row["range"], row["date"]))),
        )
        for rule, row in choices:
            selected.append(
                {
                    "quarter": quarter,
                    "selection_rule": rule,
                    "date": row["date"].isoformat(),
                    "day_index": row["day_index"],
                    "daily_mean_tCO2_per_MWh": f"{row['mean']:.8f}",
                    "intraday_range_tCO2_per_MWh": f"{row['range']:.8f}",
                    "daily_min_tCO2_per_MWh": f"{row['minimum']:.8f}",
                    "daily_max_tCO2_per_MWh": f"{row['maximum']:.8f}",
                    "result_blind_selection": 1,
                }
            )
    return selected


def artifact_hashes(output: Path, names: list[str]) -> dict[str, str]:
    return {name: sha256_path(output / name) for name in names}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.output.exists():
        raise SystemExit(f"refuse to overwrite existing output: {args.output}")
    temporary = args.output.with_name(f".{args.output.name}.tmp-{os.getpid()}")
    if temporary.exists():
        raise SystemExit(f"temporary output already exists: {temporary}")
    temporary.mkdir(parents=True)

    failures: list[str] = []
    sources, source_failures = source_checks()
    failures.extend(source_failures)
    figshare, metadata_failures = verify_figshare_metadata()
    failures.extend(metadata_failures)
    year_rows, summary_rows, values_2025, workbook_failures, workbook_warnings = audit_workbook()
    failures.extend(workbook_failures)
    slots, slot_failures = half_hour_rows(values_2025)
    failures.extend(slot_failures)
    selected = select_shanghai_days(values_2025)
    if len(selected) != 12 or len({row["date"] for row in selected}) != 12:
        failures.append("Shanghai pre-registered day selection is not 12 distinct days")

    write_csv(temporary / "raw_runs.csv", year_rows)
    write_csv(temporary / "annual_region_summary.csv", summary_rows)
    write_csv(temporary / "tvci_2025_48slot_wide.csv", slots)
    write_csv(temporary / "selected_days_shanghai_2025.csv", selected)
    status = (
        "PASS_CHINA_TVCI_2025_SOURCE_GATE_WITH_FUTURE_YEAR_NULLS_RECORDED"
        if not failures
        else "HALT_CHINA_TVCI_2025_SOURCE_GATE"
    )
    metadata = {
        "schema_version": "resetp.china-tvci-source-gate.v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "source_nature": "31-province hourly simulation/projection; not official, observed, real-time, or marginal emissions",
        "source_units": "tCO2/MWh, numerically equal to kgCO2/kWh",
        "figshare": figshare,
        "sources": sources,
        "workbook_years": list(EXPECTED_YEARS),
        "region_series_count": len(EXPECTED_COLUMNS) - 1,
        "s1_2025_hour_count": len(values_2025),
        "derived_half_hour_row_count": len(slots),
        "half_hour_conversion": "each hourly value is copied to two consecutive 30-minute slots; no interpolation",
        "day_selection": "four calendar quarters; closest-to-median daily mean, maximum intraday range, minimum intraday range; earliest date breaks ties",
        "selected_day_count": len(selected),
        "warnings": workbook_warnings,
        "authorized_scope": "S1-2025 only; 2055 and 2060 are excluded because the source workbook contains null cells",
        "search_evaluations": 0,
        "python": sys.version,
        "platform": platform.platform(),
        "script_sha256": sha256_path(Path(__file__)),
    }
    write_json(temporary / "metadata.json", metadata)
    decision = {
        "decision": status,
        "failure_count": len(failures),
        "failures": failures,
        "warning_count": len(workbook_warnings),
        "warnings": workbook_warnings,
        "search_evaluations": 0,
        "authorizes": (
            [
                "use frozen S1-2025 projection as a China TVCI scenario",
                "use generated 48-slot input",
                "use result-blind Shanghai representative days",
            ]
            if not failures
            else []
        ),
        "does_not_authorize": [
            "claiming official, historical observed, real-time, or marginal China TVCI",
            "path-search experiments before E7 closeout and ALNS G0",
            "changing dates after seeing algorithm results",
            "using the 2055 or 2060 projections without an explicit missing-data remediation contract",
        ],
    }
    write_json(temporary / "decision.json", decision)
    report = (
        "# China TVCI zero-search source gate\n\n"
        f"Decision: `{status}`.\n\n"
        f"The frozen Figshare v3 S1 workbook contains {len(EXPECTED_YEARS)} scenario years, "
        f"{len(EXPECTED_COLUMNS) - 1} regional/national series, and {len(values_2025)} hourly rows for 2025. "
        f"It was converted without interpolation to {len(slots)} half-hour rows; all daily energy-weighted integrals were checked. "
        f"Shanghai representative days were selected from the carbon curve alone before any optimization result was observed ({len(selected)} distinct days).\n\n"
        "This gate only establishes source identity, shape, completeness, numerical validity, conversion conservation, and result-blind day selection. "
        "The source remains a simulation/projection dataset and cannot be described as official, observed, real-time, or marginal emissions.\n\n"
        + (
            "Recorded non-primary-year warnings:\n"
            + "\n".join(f"- {row}" for row in workbook_warnings)
            + "\n\n"
            if workbook_warnings
            else ""
        )
        + ("Failures:\n" + "\n".join(f"- {row}" for row in failures) + "\n" if failures else "No source-gate failures were observed.\n")
    )
    write_text(temporary / "report.md", report)
    protected = [
        "raw_runs.csv",
        "annual_region_summary.csv",
        "tvci_2025_48slot_wide.csv",
        "selected_days_shanghai_2025.csv",
        "metadata.json",
        "decision.json",
        "report.md",
    ]
    write_json(
        temporary / "artifact_hashes.json",
        {
            "schema_version": "resetp.artifact-hashes.v1",
            "artifacts": artifact_hashes(temporary, protected),
        },
    )
    os.replace(temporary, args.output)
    print(json.dumps(decision, ensure_ascii=False, indent=2))
    return 0 if not failures else 2


if __name__ == "__main__":
    raise SystemExit(main())
