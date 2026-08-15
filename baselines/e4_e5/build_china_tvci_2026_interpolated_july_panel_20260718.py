#!/usr/bin/env python3
"""Build a transparent 2026 TVCI scenario and freeze the full July panel.

The upstream dataset publishes S1 hourly CEF sheets for 2025 and 2030, not
2026. This script applies piecewise-linear interpolation at each matching
hour-of-year:

    CEF_2026 = 0.8 * CEF_2025 + 0.2 * CEF_2030

No tariff, route, solver, or optimization result is read. The complete July
2026 calendar is retained; dates are not selected by outcome.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
SOURCE = REPO / "data/Carbon/中国情景/raw_20260717/CEF_data_Scenario_S1.xlsx"
OUTPUT = REPO / "baselines/e4_e5/china_tvci_2026_interpolated_july_panel_20260718"
REGIONS = ("Beijing", "Guangdong", "Chongqing")
SOURCE_YEARS = (2025, 2030)
TARGET_YEAR = 2026
WEIGHTS = {2025: 0.8, 2030: 0.2}
EXPECTED_HOURS = 8760


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_sheet(year: int) -> tuple[list[str], list[dict[str, float]]]:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(year)]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        missing = set(REGIONS) - set(headers)
        if missing:
            raise ValueError(f"{year} sheet missing regions: {sorted(missing)}")
        time_index = headers.index("Time")
        region_indices = {region: headers.index(region) for region in REGIONS}
        values: list[dict[str, float]] = []
        for expected_time, row in enumerate(rows, start=1):
            observed_time = int(row[time_index])
            if observed_time != expected_time:
                raise ValueError(
                    f"{year} time sequence mismatch: {observed_time} != {expected_time}"
                )
            record = {region: float(row[index]) for region, index in region_indices.items()}
            if not all(math.isfinite(value) for value in record.values()):
                raise ValueError(f"{year} non-finite value at hour {expected_time}")
            values.append(record)
    finally:
        workbook.close()
    if len(values) != EXPECTED_HOURS:
        raise ValueError(f"{year} expected {EXPECTED_HOURS} hours, found {len(values)}")
    return headers, values


def build_rows() -> list[dict[str, object]]:
    headers_2025, values_2025 = read_sheet(2025)
    headers_2030, values_2030 = read_sheet(2030)
    if headers_2025 != headers_2030:
        raise ValueError("2025 and 2030 sheet headers differ")

    start = datetime(TARGET_YEAR, 1, 1)
    rows: list[dict[str, object]] = []
    for hour_index, (left, right) in enumerate(
        zip(values_2025, values_2030, strict=True),
        start=1,
    ):
        timestamp = start + timedelta(hours=hour_index - 1)
        interpolated = {
            region: WEIGHTS[2025] * left[region] + WEIGHTS[2030] * right[region]
            for region in REGIONS
        }
        for region in REGIONS:
            low = min(left[region], right[region])
            high = max(left[region], right[region])
            if not low - 1e-12 <= interpolated[region] <= high + 1e-12:
                raise ValueError(f"convex-hull failure for {region} at hour {hour_index}")
        for half_hour_offset in (0, 30):
            slot_timestamp = timestamp + timedelta(minutes=half_hour_offset)
            record: dict[str, object] = {
                "date": slot_timestamp.date().isoformat(),
                "day_index": slot_timestamp.timetuple().tm_yday,
                "hourly_calendar_row": slot_timestamp.hour * 2 + half_hour_offset // 30 + 1,
                "hour_of_day": slot_timestamp.hour,
                "minute": slot_timestamp.minute,
                "source_hour_index": hour_index,
                "scenario": "S1",
                "data_nature": "piecewise_linear_interpolation_2025_2030",
            }
            for region in REGIONS:
                record[region] = f"{interpolated[region] * 1000.0:.6f}"
            rows.append(record)
    return rows


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, payload: object) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    rows = build_rows()
    if len(rows) != 365 * 48:
        raise ValueError(f"expected 17520 half-hour rows, found {len(rows)}")
    july_rows = [row for row in rows if str(row["date"]).startswith("2026-07-")]
    if len(july_rows) != 31 * 48:
        raise ValueError(f"expected 1488 July rows, found {len(july_rows)}")

    full_path = OUTPUT / "tvci_2026_interpolated_48slot_wide.csv"
    july_path = OUTPUT / "tvci_2026_july_full_month_48slot_wide.csv"
    write_csv(full_path, rows)
    write_csv(july_path, july_rows)

    source_hash = sha256(SOURCE)
    metadata = {
        "schema": "resetp.china-tvci-2026-interpolated.v1",
        "source": str(SOURCE.relative_to(REPO)),
        "source_sha256": source_hash,
        "source_dataset_doi": "10.6084/m9.figshare.28953545.v3",
        "source_paper_doi": "10.1038/s41597-026-07272-6",
        "source_scenario": "S1 baseline",
        "published_source_years": list(SOURCE_YEARS),
        "target_year": TARGET_YEAR,
        "interpolation_formula": "CEF_2026[h] = 0.8*CEF_2025[h] + 0.2*CEF_2030[h]",
        "interpolation_scope": "matching province and matching hour-of-year",
        "regions": list(REGIONS),
        "source_unit": "tCO2/MWh",
        "output_unit": "gCO2/kWh",
        "hourly_to_half_hour_rule": "copy each hourly average to two adjacent slots",
        "full_year_rows": len(rows),
        "july_rows": len(july_rows),
        "july_days": 31,
        "timezone_semantics": "Asia/Shanghai local calendar slots; not observed timestamps",
        "search_evaluations": 0,
        "optimization_results_read": False,
        "status": "PASS_CONSTRUCTED_2026_METHOD_ARCHIVE_NOT_FORMAL",
    }
    write_json(OUTPUT / "metadata.json", metadata)

    raw_rows = [
        {
            "check": "source_year_2025_hours",
            "observed": EXPECTED_HOURS,
            "expected": EXPECTED_HOURS,
            "status": "PASS",
        },
        {
            "check": "source_year_2030_hours",
            "observed": EXPECTED_HOURS,
            "expected": EXPECTED_HOURS,
            "status": "PASS",
        },
        {
            "check": "target_half_hour_rows",
            "observed": len(rows),
            "expected": 365 * 48,
            "status": "PASS",
        },
        {
            "check": "july_full_month_rows",
            "observed": len(july_rows),
            "expected": 31 * 48,
            "status": "PASS",
        },
        {
            "check": "optimization_results_read",
            "observed": 0,
            "expected": 0,
            "status": "PASS",
        },
    ]
    write_csv(OUTPUT / "raw_runs.csv", raw_rows)

    decision = {
        "decision": "PASS_CONSTRUCTED_2026_METHOD_ARCHIVE_NOT_FORMAL",
        "authorizes": [
            "audit the transparent interpolation method as an archived alternative",
        ],
        "does_not_authorize": [
            "use in the formal China E1-E7 experiment chain",
            "claiming the source directly published a 2026 sheet",
            "claiming observed, real-time, official, or marginal carbon intensity",
            "formal search before July 2026 tariff and depot contract closure",
            "dropping July dates after viewing optimization results",
        ],
        "search_evaluations": 0,
    }
    write_json(OUTPUT / "decision.json", decision)

    report = """# 中国 TVCI 2026 插值与 7 月全月面板

判定：`PASS_CONSTRUCTED_2026_METHOD_ARCHIVE_NOT_FORMAL`。

原始 Figshare S1 工作簿只发布 2025、2030、2035……2060 八个年份，没有
2026 页。本包在相同省份、相同年内小时序号上使用
`CEF_2026 = 0.8×CEF_2025 + 0.2×CEF_2030`，并把每个逐小时平均值复制为两个
半小时槽。输出是透明构造的 2026 情景，不是作者直接发布的 2026 数据，也不是
实测、实时、官方或边际排放因子。

本包保留 2026 年 7 月 1 日至 31 日全部日期，共 1488 个半小时槽。用户已裁决：
原始数据没有 2026 页即回退 2025，因此本包只作方法档案，禁止进入正式 E1–E7。

本包没有读取价格、路线、求解器或优化结果，`search_evaluations=0`。
"""
    (OUTPUT / "report.md").write_text(report, encoding="utf-8")

    artifacts = {}
    for path in sorted(OUTPUT.iterdir()):
        if path.is_file() and path.name != "artifact_hashes.json":
            artifacts[path.name] = sha256(path)
    write_json(
        OUTPUT / "artifact_hashes.json",
        {
            "hash_algorithm": "SHA-256",
            "source_sha256": source_hash,
            "artifacts": artifacts,
        },
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
